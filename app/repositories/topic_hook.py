import requests
import re
import google.generativeai as genai
from fpdf import FPDF
import os
import tempfile
import gdown
from docx import Document
from fastapi import HTTPException
from app.core.settings import settings

# Configure Gemini API
GEMINI_API_KEY = "AIzaSyBf6XLpCYFTdVG5p7YiouYxEpkGAKYqmJQ"
genai.configure(api_key=GEMINI_API_KEY)

def calling_gemini(prompt):
    """
    Calls the Gemini API to generate content based on the given prompt.
    """
    model = genai.GenerativeModel("gemini-1.5-flash")
    response = model.generate_content(prompt)
    return format_gemini_output(response.text)

def format_gemini_output(text):
    """
    Formats the output from Gemini to clean unnecessary characters while keeping line breaks.
    """
    text = re.sub(r"\*+", "", text)  # Remove asterisks
    text = re.sub(r"# +", "", text)  # Remove hashtags
    text = re.sub(r"\n{2,}", "\n", text)  # Remove excessive newlines
    return text.strip()

def fetch_lesson_plan(file_url):
    """
    Downloads a lesson plan file using gdown and processes it as a .docx file.
    """
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as temp_file:
        temp_file_path = temp_file.name

    # Use gdown to download the file with fuzzy matching
    try:
        gdown.download(url=file_url, output=temp_file_path, fuzzy=True, quiet=False)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to download file with gdown: {str(e)}")

    # Verify the file is valid
    if not os.path.exists(temp_file_path) or os.path.getsize(temp_file_path) == 0:
        raise HTTPException(status_code=400, detail="Downloaded file is empty or missing.")

    # Process the .docx file
    try:
        doc = Document(temp_file_path)
    except Exception as e:
        raise HTTPException(
            status_code=400, detail=f"Failed to process the file as a .docx: {str(e)}"
        )
    finally:
        os.remove(temp_file_path)  # Clean up temporary file

    # Extract text content from the .docx file
    return "\n".join([paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()])

def generate_dynamic_prompt(topic, audience, hook_style, learning_objective, duration):
    """
    Creates a structured prompt for generating a lesson introduction.
    """
    prompt_template = f"""
    Develop a lesson introduction (lesson hook) for {topic}, tailored for {audience}.
    
    - Provide an attention-grabbing {hook_style} to engage students at the beginning.
    - Explain the key learning objectives and the best strategies to ensure comprehension ({learning_objective}).
    - The introduction should last ({duration}).
    
    Focus on the lesson structure and execution, not the exact wording of the delivery.
    Keep the answer short and easy to understand.
    """
    return prompt_template.strip()

def generate_lesson_intro(topic, audience, hook_style, learning_objective, duration, lesson_plan_url=None):
    """
    Generates an engaging lesson introduction using Google's Gemini API, optionally integrating a lesson plan.
    """
    lesson_plan_content = ""
    if lesson_plan_url:
        lesson_plan_content = fetch_lesson_plan(lesson_plan_url)
    
    prompt = generate_dynamic_prompt(topic, audience, hook_style, learning_objective, duration)
    
    if lesson_plan_content:
        prompt += f"\n\nIncorporate the following lesson plan into the introduction:\n{lesson_plan_content}"
    
    content = calling_gemini(prompt)
    return content

def clean_text(text):
    """
    Replaces special characters with equivalent ASCII characters to avoid encoding issues.
    """
    replacements = {
        "–": "-",  # En dash
        "—": "-",  # Em dash
        "“": '"',  # Left double quote
        "”": '"',  # Right double quote
        "‘": "'",  # Left single quote
        "’": "'",  # Right single quote
        "•": "-",  # Bullet point
    }
    for char, replacement in replacements.items():
        text = text.replace(char, replacement)
    return text

class LessonPDF(FPDF):
    """
    A custom PDF class for creating lesson introduction PDFs.
    """
    def header(self):
        self.set_font("Arial", "B", 16)
        self.cell(0, 10, "Lesson Introduction", ln=True, align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", "I", 10)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")

def create_topichook_pdf(lesson_intro, filename="lesson_document.pdf"):
    """
    Creates a PDF document from the generated lesson introduction.
    """
    pdf = LessonPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    cleaned_text = clean_text(lesson_intro)  # Sanitize text for encoding compatibility
    pdf.multi_cell(0, 10, cleaned_text)
    
    full_path = os.path.join(settings.media_dir_static, filename)  # Save file in media directory
    pdf.output(full_path, "F")
    
    return full_path
